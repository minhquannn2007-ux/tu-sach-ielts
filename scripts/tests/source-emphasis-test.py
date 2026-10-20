import importlib.util,pathlib,unittest
R=pathlib.Path(__file__).resolve().parents[2];s=importlib.util.spec_from_file_location('em',R/'scripts/source-emphasis.py');em=importlib.util.module_from_spec(s);s.loader.exec_module(em)
class EmphasisTest(unittest.TestCase):
 def test_repeated_phrase_uses_context(self):
  text='In the first study Homo sapiens appeared in the title. In the second study Homo sapiens appeared in the chart.';a=text.index('Homo');flags=[a<=i<a+12 for i in range(len(text))];ps=[{'text':text}];em.apply(ps,text,flags);self.assertEqual(ps[0]['italicRanges'],[{'start':a,'end':a+12,'quote':'Homo sapiens'}])
 def test_no_fuzzy_rewrite(self):
  text='This is an entirely unrelated passage with different words and no source match.';p={'text':text};em.apply([p],'The original document contains a different and unrelated subject with several italicized statements.',[True]*101);self.assertNotIn('italicRanges',p);self.assertEqual(p['text'],text)
 def test_word_inherited_style(self):
  from docx import Document
  doc=Document();doc.styles['Normal'].font.italic=True;p=doc.add_paragraph('First italic inherited word. ');r=p.add_run('Normal override.');r.italic=False;text,flags=em.word_source(doc);self.assertTrue(all(flags[:26]));self.assertFalse(any(flags[text.index('Normal override.'):]))
 def test_misnamed_pdf_font(self):
  from unittest.mock import patch
  from types import SimpleNamespace
  class Stream:
   def get_object(self):return self
   def get_data(self):return b'embedded font'
  class Descriptor(dict):
   def get_object(self):return self
  font={'/BaseFont':'ABC+OptimaBoldItalic','/FontDescriptor':Descriptor({'/FontFile2':Stream()})}
  class Page:
   def extract_text(self,visitor_text):visitor_text('Normal body',None,None,font,12)
  fake_doc=SimpleNamespace(tobytes=lambda:b'pdf')
  with patch('pypdf.PdfReader',return_value=SimpleNamespace(pages=[Page()])),patch('fontTools.ttLib.TTFont',return_value={'head':SimpleNamespace(macStyle=0),'OS/2':SimpleNamespace(fsSelection=64)}):
   text,flags=em.pdf_source(fake_doc)
  self.assertEqual(text.strip(),'Normal body');self.assertFalse(any(flags))
 def test_utf16_offsets(self):
  text='A long sentence mentions 🧠 Homo sapiens in the discussion of human origins.';a=text.index('Homo');r=em.align(text,em.source_index(text,[a<=i<a+12 for i in range(len(text))]));self.assertEqual(r['ranges'],[{'start':a+1,'end':a+13,'quote':'Homo sapiens'}])
if __name__=='__main__':unittest.main()
